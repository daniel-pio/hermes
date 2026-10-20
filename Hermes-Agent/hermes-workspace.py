"""Scoped artifact tools. No shell, arbitrary code execution or global file access."""
import base64
import hashlib
import hmac
import json
import os
import threading
import time
import uuid
from pathlib import Path
from tools.registry import registry

PROFILES = json.loads(Path('/opt/hermes/hermes-specialists.json').read_text())
LANGUAGE_POLICY = ('Responda sempre em português do Brasil (pt-BR), inclusive relatórios e resumos de subagentes, '
    'salvo quando o usuário solicitar outra língua. Preserve nomes próprios, código, identificadores e citações '
    'no idioma original quando necessário; explique-os em português. Não traduza sintaxe de código ou SQL.')
_scopes = {}
_lock = threading.RLock()
LIMIT = 25 * 1024 * 1024

def decode_scope(token):
    encoded, signature = token.split('.', 1)
    expected = hmac.new(os.environ['API_SERVER_KEY'].encode(), encoded.encode(), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(expected, signature):
        raise ValueError('Invalid signed scope')
    scope = json.loads(base64.urlsafe_b64decode(encoded + '=' * (-len(encoded) % 4)))
    if scope['profile'] not in PROFILES or not scope.get('user') or not 0 <= time.time() - scope['issued'] <= 7200:
        raise ValueError('Invalid scope or expired request')
    scope['session'] = str(uuid.UUID(scope['session']))
    scope['run'] = str(uuid.UUID(scope['run']))
    return scope

def scope_root(scope):
    from hermes_constants import get_hermes_home
    root = get_hermes_home().resolve() / 'exports'
    path = root / scope['session'] / scope['run']
    if path.is_symlink() or not path.resolve().is_relative_to(root):
        raise ValueError('Unsafe workspace')
    path.mkdir(parents=True, exist_ok=True)
    return path.resolve()

def bind(task_id, scope):
    if scope:
        with _lock:
            _scopes[task_id] = scope

def unbind(task_id):
    with _lock:
        _scopes.pop(task_id, None)

def scoped_path(name, task_id, *, write=False):
    with _lock:
        scope = _scopes.get(task_id)
    if not scope:
        raise ValueError('No authorized workspace for this task')
    root = scope_root(scope)
    path = root / name
    if path.is_symlink() or not path.resolve().is_relative_to(root) or Path(name).is_absolute():
        raise ValueError('Path outside authorized workspace')
    if any(part.startswith('.') for part in Path(name).parts):
        raise ValueError('Hidden paths forbidden')
    if write:
        if Path(name).parts[0] == 'inputs':
            raise ValueError('Original attachments are read-only')
        path.parent.mkdir(parents=True, exist_ok=True)
    return path

def workspace_file(args, task_id=None, **kwargs):
    try:
        action = args['action']
        path = scoped_path(args.get('name', ''), task_id, write=action in ('write', 'zip', 'xlsx'))
        if action == 'list':
            result = [str(p.relative_to(path)) for p in path.rglob('*') if p.is_file() and not p.is_symlink()][:100]
        elif action == 'write':
            data = args.get('content', '').encode('utf-8')
            if len(data) > LIMIT: raise ValueError('File exceeds 25 MB')
            path.write_bytes(data)
            result = {'name': args['name'], 'size': len(data)}
        elif action == 'read':
            if path.stat().st_size > LIMIT: raise ValueError('File exceeds 25 MB')
            if path.suffix.lower() == '.pdf':
                from pypdf import PdfReader
                reader = PdfReader(path)
                start = max(1, int(args.get('page', 1)))
                result = [{'page': i + 1, 'text': reader.pages[i].extract_text() or '[Sem texto; necessita OCR]'}
                          for i in range(start - 1, min(start + 9, len(reader.pages)))]
            elif path.suffix.lower() == '.docx':
                from docx import Document
                doc = Document(path)
                result = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])[:80000]
            elif path.suffix.lower() == '.xlsx':
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                result = {s.title: [list(row) for row in s.iter_rows(values_only=True)][:500] for s in book.worksheets}
                book.close()
            else:
                result = path.read_text()[:80000]
        elif action == 'xlsx':
            from openpyxl import Workbook
            book = Workbook()
            for row in args.get('rows', []):
                # Treat strings as data, including spreadsheet formula prefixes.
                book.active.append(["'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v for v in row])
            book.save(path)
            result = {'name': args['name']}
        elif action == 'zip':
            import zipfile
            names = args.get('files', [])
            if not names or len(names) > 100: raise ValueError('Specify 1-100 files')
            with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
                for name in names:
                    source = scoped_path(name, task_id)
                    if source == path or not source.is_file() or source.stat().st_size > LIMIT: raise ValueError('Invalid ZIP input')
                    archive.write(source, arcname=name)
            if path.stat().st_size > LIMIT: path.unlink(); raise ValueError('ZIP exceeds 25 MB')
            result = {'name': args['name']}
        else:
            raise ValueError('Unknown action')
        return json.dumps({'result': result}, ensure_ascii=False, default=str)
    except Exception as exc:
        return json.dumps({'error': str(exc)})

def workspace_pdf(args, task_id=None, **kwargs):
    try:
        from pypdf import PdfReader, PdfWriter
        inputs = args['inputs']
        if not inputs or len(inputs) > 20: raise ValueError('Specify 1-20 PDFs')
        output = scoped_path(args['output'], task_id, write=True)
        if output.suffix.lower() != '.pdf': raise ValueError('Output must be PDF')
        writer = PdfWriter()
        for item in inputs:
            source = scoped_path(item['name'], task_id)
            if source == output or source.stat().st_size > LIMIT: raise ValueError('Invalid PDF input')
            reader = PdfReader(source)
            pages = item.get('pages') or list(range(1, len(reader.pages) + 1))
            for page in pages:
                if isinstance(page, bool) or not isinstance(page, int) or not 1 <= page <= len(reader.pages): raise ValueError('Invalid page number')
                if len(writer.pages) >= 1000: raise ValueError('Maximum 1000 output pages')
                writer.add_page(reader.pages[page - 1])
        writer.write(output)
        if output.stat().st_size > LIMIT: output.unlink(); raise ValueError('PDF exceeds 25 MB')
        count = len(PdfReader(output).pages)
        return json.dumps({'name': args['output'], 'pages': count, 'verified': count == len(writer.pages)})
    except Exception as exc:
        return json.dumps({'error': str(exc)})

registry.register(name='workspace_file', toolset='hermes_workspace', handler=workspace_file,
    schema={'name': 'workspace_file', 'description': 'Read/list scoped attachments or write complete output files, XLSX rows, and ZIP projects. Relative paths only. PDFs read in batches of up to 10 pages. No code execution.',
    'parameters': {'type': 'object', 'properties': {'action': {'type': 'string', 'enum': ['list','read','write','zip','xlsx']}, 'name': {'type': 'string'}, 'content': {'type': 'string'}, 'page': {'type': 'integer'}, 'files': {'type': 'array', 'items': {'type': 'string'}}, 'rows': {'type': 'array', 'items': {'type': 'array', 'items': {}}}}, 'required': ['action','name']}}, check_fn=lambda: True)
registry.register(name='workspace_pdf', toolset='hermes_workspace', handler=workspace_pdf,
    schema={'name': 'workspace_pdf', 'description': 'Merge, split or reorder scoped PDFs. Each input specifies a relative name and optional ordered page numbers (1-based). Call repeatedly for separate output PDFs. Preserves originals and verifies page count.',
    'parameters': {'type': 'object', 'properties': {'output': {'type': 'string'}, 'inputs': {'type': 'array', 'items': {'type': 'object', 'properties': {'name': {'type': 'string'}, 'pages': {'type': 'array', 'items': {'type': 'integer'}}}, 'required': ['name']}}}, 'required': ['inputs','output']}}, check_fn=lambda: True)
